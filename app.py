import os
import glob
import pandas as pd
from docx import Document
from datetime import datetime
import requests
import csv
import subprocess


def replace_placeholders_in_tables(tables, placeholders_to_values):
    for table in tables:
        for row in table.rows:
            for i, cell in enumerate(row.cells[:-1]):
                header = cell.text.strip()
                next_cell = row.cells[i + 1]
                if header == 'Opérateur':
                    next_cell.text = placeholders_to_values.get(
                        "Placeholder for the Entite Beneficiaire", "")
                elif header == 'Identifiant ADVENIR':
                    next_cell.text = placeholders_to_values.get(
                        "Placeholder for Dossier Advenir numero", "")
                elif header == 'Date':
                    next_cell.text = placeholders_to_values.get(
                        "Placeholder for the date of request exécution : JJ/MM/AAAA", "")
                elif header == 'Identifiant des points de recharge':
                    evse_ids = placeholders_to_values.get("EvseIDs", [])
                    print(evse_ids)
                    if evse_ids:
                        next_cell.text = evse_ids


def fetch_from_hubject(api_url, payload, cert_path, key_path):
    response = requests.post(api_url, json=payload, cert=(
        cert_path, key_path), verify=False)
    if response.status_code == 200:
        data = response.json()
        evse_data_records = data.get(
            'EvseData', {}).get('OperatorEvseData', [])
        # Flatten the nested structure to get a list of EvseDataRecord
        all_evse_records = [
            evse_record
            for operator_data in evse_data_records
            for evse_record in operator_data.get('EvseDataRecord', [])
        ]
        return all_evse_records
    else:
        response.raise_for_status()


def write_not_found_to_csv(not_found_evses, csv_dir):
    csv_path = os.path.join(csv_dir, 'not_found_evses.csv')
    with open(csv_path, 'w', newline='', encoding='utf-8') as file:
        writer = csv.writer(file)
        writer.writerow(['Identifiant ADVENIR', 'EvseID'])
        writer.writerows(not_found_evses)


def convert_to_pdf_libreoffice(input_file, output_directory):
    try:
        libreoffice_path = '/Applications/LibreOffice.app/Contents/MacOS/soffice'

        command = [
            libreoffice_path,
            '--headless',
            '--convert-to',
            'pdf',
            '--outdir',
            output_directory,
            input_file
        ]
        subprocess.run(command, check=True)
        print(f"Converted {input_file} to PDF.")
    except subprocess.CalledProcessError as e:
        print(f"An error occurred while converting {input_file} to PDF: {e}")


script_directory = os.path.dirname(os.path.abspath(__file__))
output_directory = os.path.join(script_directory, 'output_files')
os.makedirs(output_directory, exist_ok=True)

excel_files = glob.glob(os.path.join(script_directory, '*.xlsx'))
docx_files = glob.glob(os.path.join(script_directory, '*.docx'))
cert_files = glob.glob(os.path.join(script_directory, '*.crt'))
key_files = glob.glob(os.path.join(script_directory, '*.key'))

if not excel_files or not docx_files or not cert_files or not key_files:
    raise FileNotFoundError(
        "Required files not found in the script directory.")

df = pd.read_excel(excel_files[0])
template_document = Document(docx_files[0])
cert_path = cert_files[0]
key_path = key_files[0]

api_url = "https://service.hubject.com/api/oicp/evsepull/v22/providers/DE*ICE/data-records"
payload = {
    "GeoCoordinatesResponseFormat": "Google",
    "CountryCodes": ["FRA"],
    "ProviderID": "DE*ICE"
}

evse_data_records = fetch_from_hubject(api_url, payload, cert_path, key_path)

not_found_evses = []
for index, row in df.iterrows():
    if index == 0:
        continue

    current_date = datetime.now().strftime("%d/%m/%Y")
    entite_beneficiaire = row['Entité Bénéficiaire']
    dossier_advenir_numero = row['Dossier Advenir numéro']

    evse_ids_from_excel = [str(
        row[col_name]) for col_name in df.columns if "Point de charge" in col_name and pd.notna(row[col_name])]

    not_found_evses = []
for index, row in df.iterrows():
    if index == 0:
        continue

    current_date = datetime.now().strftime("%d/%m/%Y")
    entite_beneficiaire = row['Entité Bénéficiaire']
    dossier_advenir_numero = str(row['Dossier Advenir numéro'])

    evse_ids_from_excel = [str(
        row[col_name]) for col_name in df.columns if "Point de charge" in col_name and pd.notna(row[col_name])]

    found_evse_ids = []
    for evse_id in evse_ids_from_excel:
        if any(evse_record['EvseID'] == evse_id for evse_record in evse_data_records):
            found_evse_ids.append(evse_id)
        else:
            not_found_evses.append((dossier_advenir_numero, evse_id))

    placeholders_to_values = {
        "Placeholder for the Entite Beneficiaire": entite_beneficiaire,
        "Placeholder for Dossier Advenir numero": dossier_advenir_numero,
        "Placeholder for the date of request exécution : JJ/MM/AAAA": current_date,
        "EvseIDs": found_evse_ids
    }

    if found_evse_ids:
        evse_id_list_formatted = "\n".join(
            f"- {id_}" for id_ in found_evse_ids)
        placeholders_to_values['EvseIDs'] = evse_id_list_formatted

    replace_placeholders_in_tables(
        template_document.tables, placeholders_to_values)

    updated_docx_path = os.path.join(
        output_directory, f"Certificat Interoperabilite_{dossier_advenir_numero}.docx")
    template_document.save(updated_docx_path)

    convert_to_pdf_libreoffice(updated_docx_path, output_directory)
    os.remove(updated_docx_path)
write_not_found_to_csv(not_found_evses, output_directory)

print("Processing complete. Documents have been updated and converted to PDF.")
